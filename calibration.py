import time
import serial
import os
from openpyxl import Workbook, load_workbook

# --- CONFIGURATION ---
SERIAL_PORT = "COM3"
BAUD_RATE = 9600
EXCEL_FILENAME = "uwb_anchor_calibration.xlsx"

# Change this to switch which tag you are calibrating and which Excel tab is used
TARGET_TAG = 6

def init_excel_sheet(filename, tag_num):
    """Ensures the Excel file and the specific tag's sheet exist with headers."""
    sheet_name = f"Tag {tag_num}"

    if not os.path.exists(filename):
        wb = Workbook()
        ws = wb.active
        ws.title = sheet_name
        ws.append(["True Distance", "Measured Distance"])
        wb.save(filename)
    else:
        wb = load_workbook(filename)
        if sheet_name not in wb.sheetnames:
            ws = wb.create_sheet(title=sheet_name)
            ws.append(["True Distance", "Measured Distance"])
            wb.save(filename)

def drain_serial(ser):
    time.sleep(0.05)
    while ser.in_waiting > 0:
        ser.read(ser.in_waiting)

def check_connection(ser):
    print("Checking connection (sending AT)...")
    drain_serial(ser)
    ser.write(b"AT\r\n")
    ser.flush()

    timeout = time.time() + 2.0
    while time.time() < timeout:
        if ser.in_waiting > 0:
            line = ser.readline().decode("utf-8", errors="ignore").strip()
            print(f"Check resp: {line}")
            if "+OK" in line:
                return True
        time.sleep(0.05)
    return False

def reset_esp32(ser):
    print("Sending RESET ESP32 command...")
    ser.write(b"RESET ESP32\r\n")
    ser.flush()
    time.sleep(2.0)
    drain_serial(ser)

# --- MAIN SCRIPT ---
init_excel_sheet(EXCEL_FILENAME, TARGET_TAG)

print(f"Connecting to ESP32 on {SERIAL_PORT}...")
# Timeout set to 1 so readline() doesn't block forever if connection drops,
# but the script will loop infinitely until it gets the +OK signal.
ser = serial.Serial(SERIAL_PORT, BAUD_RATE, timeout=1)
time.sleep(2)
drain_serial(ser)

if not check_connection(ser):
    print("Initial AT check failed. Sending RESET ESP32...")
    reset_esp32(ser)
    if not check_connection(ser):
        print("FAILURE: Connection validation failed after reset. ESP32 unresponsive.")
    else:
        print("Connection valid (+OK confirmed after reset).")
else:
    print("Connection valid (+OK confirmed).")

while True:
    print(f"\n--- NEW CALIBRATION CYCLE (Logging to 'Tag {TARGET_TAG}' tab) ---")
    try:
        true_distance = float(input("Enter True Distance (cm): "))
    except ValueError:
        print("Invalid number input, skipping.")
        continue

    input("Press Enter to prompt ESP32 and start the test...")

    drain_serial(ser)
    ser.write(b"SETUP ANCHOR\r\n")
    ser.flush()
    print("Test running... waiting for Data Points (No timeout).")

    session_data = []

    # Infinite loop to wait for data
    while True:
        line = ser.readline().decode("utf-8", errors="ignore").strip()

        if line:
            print(f"ESP32: {line}")

            # Parse lines exactly like "Data Point 1: 44"
            if "Data Point" in line:
                try:
                    # Splitting by ':' gets everything after the colon
                    val_str = line.split(":")[-1].strip()
                    dist = int(val_str)

                    # Save the inputted True Distance alongside the Measured Distance
                    session_data.append([true_distance, dist])
                except ValueError:
                    pass

            # Stop reading once the final confirmation message is received
            elif "Received: +OK" in line:
                break

    # Save collected data to Excel
    if session_data:
        wb = load_workbook(EXCEL_FILENAME)
        ws = wb[f"Tag {TARGET_TAG}"]
        for row in session_data:
            ws.append(row)
        wb.save(EXCEL_FILENAME)
        print(f"Test Complete. Saved {len(session_data)} data points for Tag {TARGET_TAG}.")
    else:
        print("Test Complete. WARNING: No data points were received.")

    next_coords = input("\nPress Enter to test another distance or type 'q' to quit: ")

    if next_coords.strip().lower() == "q":
        break

ser.close()
